"""Executable entry point; the hotkey agent shares the bundled runtime."""
import sys


def self_test():
    import io
    import json
    import threading
    import urllib.request
    import docx
    import pptx
    import openpyxl
    import tkinter
    from pypdf import PdfReader, PdfWriter
    from sklearn.feature_extraction.text import TfidfVectorizer
    import server
    import background_agent

    for document in (docx.Document(), pptx.Presentation(), openpyxl.Workbook()):
        output = io.BytesIO()
        document.save(output)
        assert output.tell() > 0
    pdf = PdfWriter()
    pdf.add_blank_page(width=72, height=72)
    output = io.BytesIO()
    pdf.write(output)
    output.seek(0)
    assert len(PdfReader(output).pages) == 1
    assert TfidfVectorizer().fit_transform(['hello world', 'hello application']).shape == (2, 3)
    assert tkinter.Tcl().eval('expr {1 + 1}') == '2'
    httpd = server.ThreadingHTTPServer(('127.0.0.1', 0), server.Handler)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    try:
        for route in ('/', '/chat'):
            with urllib.request.urlopen(f'http://127.0.0.1:{httpd.server_port}{route}') as response:
                assert response.status == 200 and len(response.read()) > 100
    finally:
        httpd.shutdown()
        httpd.server_close()
    print(json.dumps({'self_test': 'PASS', 'frozen': getattr(sys, 'frozen', False), 'root': server.ROOT}))


if __name__ == '__main__':
    for stream in (sys.stdout, sys.stderr):
        if stream is not None:
            stream.reconfigure(encoding='utf-8', errors='replace')
    if '--self-test' in sys.argv:
        self_test()
    elif '--background-agent' in sys.argv:
        import background_agent
        background_agent.main()
    elif '--ocr-plugin' in sys.argv:
        from plugins.ocr import ocr_bridge
        ocr_bridge.main()
    else:
        import server
        server.main()
