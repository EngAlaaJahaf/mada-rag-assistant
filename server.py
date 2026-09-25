# -*- coding: utf-8 -*-
"""
مذكّرتي — RAG محلي بدون إنترنت.
يستخرج النصوص من txt/md/docx/pptx/pdf، يجيب بالاقتباس من المحتوى،
ويولّد إجابة عربية مشروحة عبر نموذج Qwen3 المحلي (llama.cpp) عند تفعيله.
"""
import os
import re
import time
import json
import sys
import atexit
import warnings
import threading
import urllib.request
import subprocess
import queue
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse, parse_qs, unquote
import ctypes

for _s in (sys.stdout, sys.stderr):
    if _s is not None:
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

warnings.filterwarnings("ignore", message="The parameter 'token_pattern'.*")

import docx
import pptx
import openpyxl
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer

ROOT = os.path.dirname(sys.executable if getattr(sys, "frozen", False) else os.path.abspath(__file__))


def get_asset_path(rel_path):
    """يرجع مسار الملف الثابت مع إعطاء الأولوية للمجلد المحلي بجانب التطبيق،
    أو استخراجه من الحزمة المدمجة (_MEIPASS) عند التشغيل كملف تنفيذي مدمج."""
    local = os.path.join(ROOT, rel_path)
    if os.path.isfile(local):
        return local
    meipass = getattr(sys, "_MEIPASS", None)
    if meipass:
        bundled = os.path.join(meipass, rel_path)
        if os.path.isfile(bundled):
            return bundled
    return local


UPLOADS = os.path.join(ROOT, "uploads")
CHUNKS_PATH = os.path.join(ROOT, "index_data", "chunks.json")
LOCK = threading.Lock()

ALLOWED = {".txt", ".md", ".pdf", ".docx", ".pptx", ".xlsx", ".sql", ".json", ".csv", ".py", ".js", ".html", ".css", ".log"}

# ── طبقة التخزين المحلي SQLite للمحادثات والمشروعات ──
import sqlite3

DB_PATH = os.path.join(ROOT, "mada.db")
DB_LOCK = threading.Lock()

def get_db():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn

def init_db():
    with DB_LOCK:
        conn = get_db()
        with conn:
            conn.execute("""
            CREATE TABLE IF NOT EXISTS projects(
              id TEXT PRIMARY KEY,
              name TEXT NOT NULL,
              color TEXT NOT NULL DEFAULT '#10a37f',
              icon TEXT NOT NULL DEFAULT 'folder',
              created_at TEXT,
              updated_at TEXT
            );
            """)
            conn.execute("""
            CREATE TABLE IF NOT EXISTS conversations(
              id TEXT PRIMARY KEY,
              project_id TEXT REFERENCES projects(id) ON DELETE SET NULL,
              mode TEXT DEFAULT 'free',
              title TEXT DEFAULT '',
              messages TEXT DEFAULT '[]',
              pinned INTEGER DEFAULT 0,
              archived INTEGER DEFAULT 0,
              created_at TEXT,
              updated_at TEXT
            );
            """)
            conn.execute("""
            CREATE TABLE IF NOT EXISTS memory(
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              kind TEXT NOT NULL CHECK(kind IN ('term','rule')),
              subject TEXT NOT NULL,
              wrong TEXT DEFAULT '',
              note TEXT DEFAULT '',
              created_at TEXT DEFAULT (datetime('now'))
            );
            """)
        conn.close()

def list_convs(include_archived=False):
    with DB_LOCK:
        conn = get_db()
        query = "SELECT * FROM conversations"
        if not include_archived:
            query += " WHERE archived = 0"
        query += " ORDER BY pinned DESC, updated_at DESC"
        rows = conn.execute(query).fetchall()
        out = []
        for r in rows:
            try:
                msgs = json.loads(r["messages"] or "[]")
            except Exception:
                msgs = []
            out.append({
                "id": r["id"],
                "project_id": r["project_id"],
                "mode": r["mode"] or "free",
                "title": r["title"] or "",
                "messages": msgs,
                "pinned": bool(r["pinned"]),
                "archived": bool(r["archived"]),
                "created_at": r["created_at"],
                "updated_at": r["updated_at"]
            })
        conn.close()
        return out

def get_conv(cid):
    with DB_LOCK:
        conn = get_db()
        r = conn.execute("SELECT * FROM conversations WHERE id = ?", (cid,)).fetchone()
        conn.close()
        if not r:
            return None
        try:
            msgs = json.loads(r["messages"] or "[]")
        except Exception:
            msgs = []
        return {
            "id": r["id"],
            "project_id": r["project_id"],
            "mode": r["mode"] or "free",
            "title": r["title"] or "",
            "messages": msgs,
            "pinned": bool(r["pinned"]),
            "archived": bool(r["archived"]),
            "created_at": r["created_at"],
            "updated_at": r["updated_at"]
        }

def upsert_conv(chat):
    if not isinstance(chat, dict):
        return False
    cid = str(chat.get("id") or "").strip()
    if not cid:
        return False
    
    pid = chat.get("project_id")
    if isinstance(pid, dict):
        pid = pid.get("id")
    pid = str(pid).strip() if (pid is not None and str(pid).strip()) else None

    mode = chat.get("mode")
    if isinstance(mode, dict):
        mode = mode.get("id") or mode.get("mode")
    mode = str(mode or "free").strip()

    title = chat.get("title")
    if isinstance(title, dict):
        title = title.get("title") or ""
    title = str(title or "")

    msgs_raw = chat.get("messages")
    if isinstance(msgs_raw, list):
        msgs_json = json.dumps(msgs_raw, ensure_ascii=False)
    elif isinstance(msgs_raw, str):
        msgs_json = msgs_raw
    else:
        msgs_json = "[]"

    pinned = 1 if chat.get("pinned") else 0
    archived = 1 if chat.get("archived") else 0
    now = time.strftime("%Y-%m-%d %H:%M:%S")

    with DB_LOCK:
        conn = get_db()
        with conn:
            if pid:
                p_exists = conn.execute("SELECT 1 FROM projects WHERE id = ?", (pid,)).fetchone()
                if not p_exists:
                    pid = None
            existing = conn.execute("SELECT created_at, pinned, archived FROM conversations WHERE id = ?", (cid,)).fetchone()
            if existing:
                created_at = existing["created_at"] or now
                if "pinned" not in chat:
                    pinned = existing["pinned"]
                if "archived" not in chat:
                    archived = existing["archived"]
                conn.execute("""
                UPDATE conversations
                SET project_id = ?, mode = ?, title = ?, messages = ?, pinned = ?, archived = ?, updated_at = ?
                WHERE id = ?
                """, (pid, mode, title, msgs_json, pinned, archived, now, cid))
            else:
                created_at = str(chat.get("created_at") or now)
                conn.execute("""
                INSERT INTO conversations (id, project_id, mode, title, messages, pinned, archived, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (cid, pid, mode, title, msgs_json, pinned, archived, created_at, now))
        conn.close()
    return True

def delete_conv(cid):
    if not cid:
        return False
    cid = str(cid).strip()
    with DB_LOCK:
        conn = get_db()
        with conn:
            conn.execute("DELETE FROM conversations WHERE id = ?", (cid,))
        conn.close()
    return True

def move_conv(cid, pid):
    cid = str(cid or "").strip()
    if not cid:
        return False
    if isinstance(pid, dict):
        pid = pid.get("id")
    pid = str(pid).strip() if (pid is not None and str(pid).strip()) else None
    now = time.strftime("%Y-%m-%d %H:%M:%S")
    with DB_LOCK:
        conn = get_db()
        with conn:
            conn.execute("UPDATE conversations SET project_id = ?, updated_at = ? WHERE id = ?", (pid, now, cid))
        conn.close()
    return True

def pin_conv(cid, pinned):
    cid = str(cid or "").strip()
    if not cid:
        return False
    now = time.strftime("%Y-%m-%d %H:%M:%S")
    with DB_LOCK:
        conn = get_db()
        with conn:
            conn.execute("UPDATE conversations SET pinned = ?, updated_at = ? WHERE id = ?", (1 if pinned else 0, now, cid))
        conn.close()
    return True

def archive_conv(cid, archived):
    cid = str(cid or "").strip()
    if not cid:
        return False
    now = time.strftime("%Y-%m-%d %H:%M:%S")
    with DB_LOCK:
        conn = get_db()
        with conn:
            conn.execute("UPDATE conversations SET archived = ?, updated_at = ? WHERE id = ?", (1 if archived else 0, now, cid))
        conn.close()
    return True

def clear_convs():
    with DB_LOCK:
        conn = get_db()
        with conn:
            conn.execute("DELETE FROM conversations")
        conn.close()
    return True

def list_projects():
    with DB_LOCK:
        conn = get_db()
        rows = conn.execute("SELECT * FROM projects ORDER BY created_at ASC").fetchall()
        out = []
        for r in rows:
            c_cnt = conn.execute("SELECT COUNT(*) as cnt FROM conversations WHERE project_id = ? AND archived = 0", (r["id"],)).fetchone()["cnt"]
            out.append({
                "id": r["id"],
                "name": r["name"],
                "color": r["color"] or "#10a37f",
                "icon": r["icon"] or "folder",
                "count": c_cnt,
                "created_at": r["created_at"],
                "updated_at": r["updated_at"]
            })
        conn.close()
        return out

def upsert_project(p):
    if not isinstance(p, dict):
        return None
    pid = str(p.get("id") or ("p_" + str(int(time.time() * 1000)))).strip()
    name = str(p.get("name") or "مشروع جديد").strip()
    color = str(p.get("color") or "#10a37f").strip()
    icon = str(p.get("icon") or "folder").strip()
    now = time.strftime("%Y-%m-%d %H:%M:%S")

    with DB_LOCK:
        conn = get_db()
        with conn:
            existing = conn.execute("SELECT created_at FROM projects WHERE id = ?", (pid,)).fetchone()
            if existing:
                conn.execute("""
                UPDATE projects SET name = ?, color = ?, icon = ?, updated_at = ?
                WHERE id = ?
                """, (name, color, icon, now, pid))
            else:
                conn.execute("""
                INSERT INTO projects (id, name, color, icon, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?)
                """, (pid, name, color, icon, now, now))
        conn.close()
    return {"id": pid, "name": name, "color": color, "icon": icon}

def delete_project(pid):
    with DB_LOCK:
        conn = get_db()
        with conn:
            conn.execute("DELETE FROM projects WHERE id = ?", (pid,))
        conn.close()
    return True

# ── إعداد الترجمة المحلية (إنجليزي→عربي) عبر نموذج LibreTranslate على جهازك ──
VENV_PY = r"D:\projects\LibreTranslate-1.9.6\venv\Scripts\python.exe"
BRIDGE_SCRIPT = os.path.join(ROOT, "translate_bridge.py")
TRANSLATE_MAX_RESULTS = 3
TRANSLATE_MAX_CHARS = 900

_bridge_lock = threading.Lock()
_bridge_proc = None
_bridge_id = 0


def translation_available():
    return os.path.isfile(VENV_PY) and os.path.isfile(BRIDGE_SCRIPT)


def ensure_bridge():
    global _bridge_proc
    with _bridge_lock:
        if _bridge_proc and _bridge_proc.poll() is None:
            return True
        try:
            _bridge_proc = subprocess.Popen(
                [VENV_PY, BRIDGE_SCRIPT],
                stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                text=True, encoding="utf-8", errors="replace",
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
            _procs.append(_bridge_proc)
            return True
        except Exception:
            _bridge_proc = None
            return False


def translate_text(text, direction="en2ar"):
    if not translation_available():
        return None
    global _bridge_id
    if not ensure_bridge():
        return None
    _bridge_id += 1
    mid = _bridge_id
    msg = json.dumps({"id": mid, "text": text[:TRANSLATE_MAX_CHARS], "dir": direction}, ensure_ascii=False) + "\n"
    with _bridge_lock:
        try:
            _bridge_proc.stdin.write(msg)
            _bridge_proc.stdin.flush()
            for _ in range(2):
                line = _bridge_proc.stdout.readline()
                if not line:
                    return None
                try:
                    out = json.loads(line)
                except ValueError:
                    continue
                if out.get("id") == mid:
                    ar = out.get("ar") if out.get("ok") else None
                    if ar:
                        ar = apply_memory_terms(ar)
                    return ar
        except Exception:
            return None
    return None


def translate_results(results, n=TRANSLATE_MAX_RESULTS):
    for r in results[:n]:
        ar = translate_text(r["text"])
        if ar and ar.strip():
            r["ar"] = ar.strip()
    return results


# ── ذاكرة التعلم: قواعد ومصطلحات يتعلمها النموذج من تصحيحات المستخدم ──
def list_memory(kind=None):
    with DB_LOCK:
        conn = get_db()
        if kind:
            rows = conn.execute("SELECT * FROM memory WHERE kind=? ORDER BY id DESC", (kind,)).fetchall()
        else:
            rows = conn.execute("SELECT * FROM memory ORDER BY id DESC").fetchall()
        conn.close()
        return [{
            "id": r["id"],
            "kind": r["kind"],
            "subject": r["subject"],
            "wrong": r["wrong"] or "",
            "note": r["note"] or "",
            "created_at": r["created_at"] or "",
        } for r in rows]


def add_memory(kind, subject, wrong="", note=""):
    with DB_LOCK:
        conn = get_db()
        with conn:
            cur = conn.execute(
                "INSERT INTO memory(kind, subject, wrong, note) VALUES(?,?,?,?)",
                (kind, subject, wrong or "", note or ""),
            )
            new_id = cur.lastrowid
        conn.close()
        return new_id


def delete_memory(mid):
    with DB_LOCK:
        conn = get_db()
        with conn:
            conn.execute("DELETE FROM memory WHERE id=?", (mid,))
        conn.close()


def update_memory(mid, kind, subject, wrong="", note=""):
    with DB_LOCK:
        conn = get_db()
        with conn:
            conn.execute(
                "UPDATE memory SET kind=?, subject=?, wrong=?, note=? WHERE id=?",
                (kind, subject, wrong or "", note or "", mid),
            )
        conn.close()


def memory_rules_block():
    """نص القواعد المتراكمة لِحَقنِه في ترويسة النموذج."""
    rows = list_memory()
    if not rows:
        return ""
    lines = ["[قواعد تعلّمتها من المستخدم سابقاً والتزم بها في ردودك العربية]:", ""]
    for r in rows:
        if r["kind"] == "term":
            lines.append("• المصطلح «%s» هو الصواب، ولا تكتبه «%s»." % (r["subject"], r["wrong"] or "…"))
        else:
            lines.append("• %s" % r["subject"])
        if r["note"]:
            lines[-1] += " ملاحظة: %s" % r["note"]
    return "\n".join(lines)


def apply_memory_terms(text):
    """استبدال حتمي للصيغ الخاطئة المعروفة (مصطلحات) في أي نص عربي ناتج."""
    if not text:
        return text
    rows = list_memory("term")
    for r in rows:
        w = (r.get("wrong") or "").strip()
        s = (r.get("subject") or "").strip()
        if w and s and w != s:
            try:
                text = re.sub(re.escape(w), s, text, flags=re.IGNORECASE)
            except Exception:
                pass
    return text


# ── التوليد المحلي (نموذج Qwen3 عبر llama.cpp على جهازك) ──
# المسارات الافتراضية (يمكن تغييرها من الواجهة وتُحفظ في model_config.json)
LLM_EXE_DEFAULT = r"D:\jan\llamacpp\backends\b9967\win-avx-cuda-cu12.0-x64\build\bin\llama-server.exe"
LLM_MODEL_DEFAULT = r"D:\jan\llamacpp\models\Jan-v3.5-4B-Q4_K_XL\model.gguf"
LLM_EXE = LLM_EXE_DEFAULT
LLM_MODEL = LLM_MODEL_DEFAULT
LLM_HOST = "127.0.0.1"
LLM_PORT = 8081
LLM_MAX_TOKENS = 450
LLM_CTX = 2048
LLM_NGL = 0
LLM_THREADS = 0

_gen_lock = threading.Lock()
_procs = []
_llm_proc = None
_llm_start_lock = threading.RLock()
_llm_error = ""
LLM_LOG_PATH = os.path.join(ROOT, "llama-server.log")
_agent_proc = None

BG_PATH = os.path.join(ROOT, "bg_config.json")
MODEL_CONFIG_PATH = os.path.join(ROOT, "model_config.json")

DEFAULT_BG = {
    "enabled": True,
    "hotkey": "Ctrl+Alt+Space",
    "reopen_hotkey": "Ctrl+Alt+R",
    "length": "short",
    "custom_command": "",
    "duration": 3,
    "position": "right",
    "color": "dark",
    "size": "medium",
    "custom_width": None,
    "custom_height": None,
    "custom_x": None,
    "custom_y": None,
    "opacity": 1.0,
    "saveReplies": True,
}

DEFAULT_MODEL = {
    "spec": "auto",
    "exe": LLM_EXE_DEFAULT,
    "model": LLM_MODEL_DEFAULT,
}

BG_BOOLS = ("enabled", "saveReplies")
BG_CHOICES = {
    "length": ("short", "medium", "detailed"),
    "position": ("left", "center", "right"),
    "color": ("dark", "light"),
    "size": ("small", "medium", "wide", "custom"),
}


def load_bg_config():
    out = dict(DEFAULT_BG)
    try:
        if os.path.isfile(BG_PATH):
            with open(BG_PATH, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, dict):
                out.update(data)
        else:
            try:
                save_bg_config(out)
            except Exception:
                pass
    except Exception:
        pass
    return out


def save_bg_config(cfg):
    with open(BG_PATH, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)
    return cfg


def sanitize_bg(data):
    if not isinstance(data, dict):
        return load_bg_config()
    out = load_bg_config()
    for k in BG_BOOLS:
        if k in data:
            out[k] = bool(data[k])
    d = data.get("duration")
    if d is not None:
        try:
            val = float(d)
            clamped = max(0.1, min(30.0, val))
            out["duration"] = round(clamped, 2) if clamped % 1 != 0 else int(clamped)
        except (ValueError, TypeError):
            pass
    for k, allowed in BG_CHOICES.items():
        if k in data and str(data[k]) in allowed:
            out[k] = str(data[k])
    for dim_k, (min_v, max_v) in (("custom_width", (240, 1920)), ("custom_height", (140, 1200))):
        if dim_k in data:
            v = data.get(dim_k)
            if v is None:
                out[dim_k] = None
            else:
                try:
                    val = int(v)
                    out[dim_k] = max(min_v, min(max_v, val))
                except (ValueError, TypeError):
                    pass
    # حفظ موضع النافذة (x, y) في حدود معقولة
    pos_explicitly_sent = "position" in data
    coords_explicitly_sent = ("custom_x" in data or "custom_y" in data)
    if pos_explicitly_sent and not coords_explicitly_sent:
        out["custom_x"] = None
        out["custom_y"] = None
    else:
        for pos_k, (min_v, max_v) in (("custom_x", (-200, 3840)), ("custom_y", (-200, 2160))):
            if pos_k in data:
                v = data.get(pos_k)
                if v is None:
                    out[pos_k] = None
                else:
                    try:
                        val = int(v)
                        out[pos_k] = max(min_v, min(max_v, val))
                    except (ValueError, TypeError):
                        pass
    # حفظ شفافية النافذة في نطاق 0.2–1.0
    if "opacity" in data:
        v = data.get("opacity")
        if v is None:
            out["opacity"] = 1.0
        else:
            try:
                val = float(v)
                out["opacity"] = round(max(0.2, min(1.0, val)), 2)
            except (ValueError, TypeError):
                pass
    if out.get("size") != "custom" and "size" in data and data["size"] in ("small", "medium", "wide"):
        out["custom_width"] = None
        out["custom_height"] = None
    for k in ("hotkey", "reopen_hotkey"):
        v = str(data.get(k) or "").strip()
        if v and len(v) <= 32 and re.fullmatch(r"[A-Za-z0-9+ ]+", v):
            out[k] = v
    cc = str(data.get("custom_command") or "").strip()
    out["custom_command"] = cc[:1000]
    return out


MODEL_MTOKENS = {"auto": 600, "low": 300, "medium": 450, "full": 600}
SPEC_PARAMS = {
    "auto": {"ctx": 8192, "ngl": 99},
    "low": {"ctx": 4096, "ngl": 0},
    "medium": {"ctx": 4096, "ngl": 24},
    "full": {"ctx": 8192, "ngl": 99},
}


def detect_system_hardware():
    """يكتشف مواصفات الجهاز الحالي (الرام، الأنوية، كارت الشاشة NVIDIA وحجم الذاكرة)."""
    hw = {
        "gpu_name": None,
        "gpu_vram_mb": 0,
        "has_nvidia": False,
        "ram_total_gb": 8.0,
        "ram_free_gb": 4.0,
        "cpu_cores": os.cpu_count() or 4,
    }
    try:
        class _MEM(ctypes.Structure):
            _fields_ = [
                ("dwLength", ctypes.c_ulong), ("dwMemoryLoad", ctypes.c_ulong),
                ("ullTotalPhys", ctypes.c_ulonglong), ("ullAvailPhys", ctypes.c_ulonglong),
                ("ullTotalPageFile", ctypes.c_ulonglong), ("ullAvailPageFile", ctypes.c_ulonglong),
                ("ullTotalVirtual", ctypes.c_ulonglong), ("ullAvailVirtual", ctypes.c_ulonglong),
                ("sullAvailExtendedVirtual", ctypes.c_ulonglong),
            ]
        m = _MEM()
        m.dwLength = ctypes.sizeof(_MEM)
        if ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(m)):
            hw["ram_total_gb"] = round(m.ullTotalPhys / (1024**3), 1)
            hw["ram_free_gb"] = round(m.ullAvailPhys / (1024**3), 1)
    except Exception:
        pass

    try:
        r = subprocess.run(
            ["nvidia-smi", "--query-gpu=name,memory.total", "--format=csv,noheader,nounits"],
            capture_output=True, text=True, timeout=2, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        if r.returncode == 0 and r.stdout.strip():
            line = r.stdout.strip().splitlines()[0]
            parts = [p.strip() for p in line.split(",")]
            hw["gpu_name"] = parts[0]
            hw["gpu_vram_mb"] = int(parts[1]) if len(parts) > 1 and parts[1].isdigit() else 0
            hw["has_nvidia"] = True
    except Exception:
        pass

    return hw


def derive_auto_specs(hw=None, exe_path=""):
    """يشتق أفضل إعدادات أداء ديناميكية (سياق، طبقات GPU، خيوط) وفق عتاد الجهاز."""
    if hw is None:
        hw = detect_system_hardware()
    vram = hw.get("gpu_vram_mb", 0)
    ram = hw.get("ram_total_gb", 8.0)
    cores = hw.get("cpu_cores", 4)
    is_cuda_bin = "cuda" in str(exe_path).lower() if exe_path else hw["has_nvidia"]

    if hw["has_nvidia"] and is_cuda_bin and vram >= 5500:
        return {
            "effective_spec": "full",
            "ngl": 99,
            "ctx": 8192,
            "threads": min(8, max(4, cores - 2)),
            "max_tokens": 600,
            "reason": f"كارت شاشة قوي ({hw['gpu_name']} - {round(vram/1024, 1)}GB VRAM): تسريع كامل على GPU وسياق 8192 بأعلى سرعة.",
        }
    elif hw["has_nvidia"] and is_cuda_bin and vram >= 3500:
        return {
            "effective_spec": "medium",
            "ngl": 24,
            "ctx": 4096,
            "threads": min(8, max(4, cores - 2)),
            "max_tokens": 450,
            "reason": f"كارت شاشة متوسط ({hw['gpu_name']} - {round(vram/1024, 1)}GB VRAM): تسريع هجين (24 طبقة GPU) وسياق 4096 لموازنة الذاكرة.",
        }
    elif hw["has_nvidia"] and is_cuda_bin and vram >= 1800:
        return {
            "effective_spec": "low",
            "ngl": 12,
            "ctx": 4096,
            "threads": max(2, cores // 2),
            "max_tokens": 300,
            "reason": f"كارت شاشة بذاكرة محدودة ({round(vram/1024, 1)}GB VRAM): ترحيل جزئي خفيف (12 طبقة) وسياق 4096.",
        }
    else:
        rec_threads = max(2, min(8, cores - 1 if cores > 4 else max(2, cores // 2)))
        rec_ctx = 8192 if ram >= 16.0 else 4096
        return {
            "effective_spec": "cpu",
            "ngl": 0,
            "ctx": rec_ctx,
            "threads": rec_threads,
            "max_tokens": 350,
            "reason": f"تشغيل آمن على المعالج ({cores} أنوية، {ram}GB RAM): لا يتوفر GPU متوافق، تم ضبط {rec_threads} خيوط وسياق {rec_ctx}.",
        }


def _resolve_threads(spec):
    n = os.cpu_count() or 4
    if spec == "low":
        return max(2, n // 2)
    if spec == "medium":
        return max(2, n - 1)
    return 8


def load_model_config():
    out = dict(DEFAULT_MODEL)
    try:
        if os.path.isfile(MODEL_CONFIG_PATH):
            with open(MODEL_CONFIG_PATH, "r", encoding="utf-8") as f:
                data = json.load(f)
            if isinstance(data, dict):
                out.update(data)
    except Exception:
        pass
    return out


PORTABLE_DIR = os.path.join(ROOT, "portable")


def discover_portable_llm():
    """يبحث عن مشغّل النموذج والنموذج داخل مجلد portable بجوار التطبيق.

    إرجاع (exe, model) إن وُجدا فعلاً، أو (None, None). يُفضَّل المسار المطلق
    المثبَّت أصلاً (D:\\jan…) ويُستخدم تلقائياً إن وُجد؛ بينما portable حُلّة
    النسخة المستقلة المنقولة على فلاشة/جهاز آخر دون الحاجة لتعديل يدوي.
    """
    exe = os.path.join(PORTABLE_DIR, "llama", "llama-server.exe")
    if not os.path.isfile(exe):
        exe = os.path.join(PORTABLE_DIR, "bin", "llama-server.exe")
    if not os.path.isfile(exe):
        return None, None

    # النموذج: gguf واحد أو أكبر gguf ضمن مجلد models
    model = os.path.join(PORTABLE_DIR, "models", "model.gguf")
    if not os.path.isfile(model):
        cands = []
        try:
            for root, _, files in os.walk(os.path.join(PORTABLE_DIR, "models")):
                for f in files:
                    if f.lower().endswith(".gguf"):
                        p = os.path.join(root, f)
                        cands.append((os.path.getsize(p), p))
        except Exception:
            pass
        if not cands:
            return None, None
        cands.sort(reverse=True)
        model = cands[0][1]

    return exe, model


def get_llm_config():
    """يرجع إعدادات النموذج المحلية الموحّدة مع اشتقاق بارامترات الأداء الذكي وفق العتاد."""
    cfg = load_model_config()
    spec = str(cfg.get("spec") if cfg.get("spec") in SPEC_PARAMS else "auto")
    exe = str(cfg.get("exe") or LLM_EXE_DEFAULT).strip() or LLM_EXE_DEFAULT
    model = str(cfg.get("model") or LLM_MODEL_DEFAULT).strip() or LLM_MODEL_DEFAULT
    # Resolve before changing the child process working directory to its bin folder.
    exe = os.path.abspath(os.path.join(ROOT, exe))
    model = os.path.abspath(os.path.join(ROOT, model))
    # بديل تلقائي محمول: إن لم يجد الملفان في مواضعهما المُهيأة (مثلاً على فلاشة
    # بلا التثبيت على D:\) فابحث داخل مجلد portable\ بجوار التطبيق كحل مستقل.
    if not (os.path.isfile(exe) and os.path.isfile(model)):
        pexe, pmodel = discover_portable_llm()
        if pexe and pmodel:
            exe, model = pexe, pmodel

    hw = detect_system_hardware()
    auto = derive_auto_specs(hw, exe)
    if spec == "auto":
        ctx = auto["ctx"]
        ngl = auto["ngl"]
        threads = auto["threads"]
        max_tokens = auto["max_tokens"]
        effective = auto["effective_spec"]
        reason = auto["reason"]
    else:
        p = SPEC_PARAMS.get(spec, SPEC_PARAMS["full"])
        ctx = int(cfg.get("ctx") or p["ctx"])
        ngl = int(cfg.get("ngl") if cfg.get("ngl") is not None else p["ngl"])
        threads = int(cfg.get("threads") or _resolve_threads(spec))
        max_tokens = MODEL_MTOKENS.get(spec, 450)
        effective = spec
        reason = f"تم الضبط يدوياً ({spec})"

    return {
        "spec": spec,
        "effective_spec": effective,
        "exe": exe,
        "model": model,
        "host": LLM_HOST,
        "port": LLM_PORT,
        "ctx": ctx,
        "ngl": ngl,
        "threads": threads,
        "max_tokens": max_tokens,
        "hardware": hw,
        "auto_derived": auto,
        "reason": reason,
    }


def _cleanup_childs():
    stop_background_agent()
    for p in list(_procs):
        try:
            if p and p.poll() is None:
                p.terminate()
        except Exception:
            pass


atexit.register(_cleanup_childs)


def generation_available():
    c = get_llm_config()
    return os.path.isfile(c["exe"]) and os.path.isfile(c["model"])


def _llm_alive():
    try:
        with urllib.request.urlopen(f"http://{LLM_HOST}:{LLM_PORT}/v1/models", timeout=2) as r:
            return r.status == 200
    except Exception:
        return False


def stop_llm():
    global _llm_proc
    with _llm_start_lock:
        if _llm_proc is not None:
            try:
                if _llm_proc.poll() is None:
                    _llm_proc.terminate()
                    _llm_proc.wait(timeout=5)
            except subprocess.TimeoutExpired:
                _llm_proc.kill()
                _llm_proc.wait(timeout=5)
            finally:
                if _llm_proc in _procs:
                    _procs.remove(_llm_proc)
                _llm_proc = None


def ensure_llm():
    global _llm_proc, _llm_error
    if _llm_alive():
        return True
    if not generation_available():
        _llm_error = "ملف المحرك أو النموذج غير موجود؛ راجع إعدادات النموذج."
        return False
    c = get_llm_config()
    with _llm_start_lock:
        if _llm_alive():
            return True
        if _llm_proc is not None and _llm_proc.poll() is None:
            return True
        if _llm_proc in _procs:
            _procs.remove(_llm_proc)
        try:
            with open(LLM_LOG_PATH, "w", encoding="utf-8") as log:
                proc = subprocess.Popen(
                    [c["exe"], "-m", c["model"],
                     "--host", LLM_HOST, "--port", str(LLM_PORT),
                     "--ctx-size", str(c["ctx"]), "-ngl", str(c["ngl"]), "-t", str(c["threads"])],
                    cwd=os.path.dirname(c["exe"]) or None,
                    creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
                    stdout=log, stderr=subprocess.STDOUT,
                )
            _llm_proc = proc
            _llm_error = ""
            _procs.append(proc)
            return True
        except Exception as exc:
            _llm_error = "تعذر تشغيل المحرك: " + str(exc)
            return False


def generate_answer(query, results, target_file=None):
    ensure_llm()
    for _ in range(80):
        if _llm_alive():
            break
        time.sleep(0.5)
    if not _llm_alive():
        return None
    lines = []
    for r in results[:8]:
        t = r.get("text", "").strip()[:800]
        a = r.get("ar", "").strip()[:500]
        lines.append("[%s]\n%s" % (r.get("file", ""), t))
        if a:
            lines.append("بالعربية: " + a)
    chunks = "\n\n".join(lines)
    if target_file:
        system = (
            f"أنت مساعد مذاكرة يجيب بالعربية الفصحى اعتماداً فقط وحصرياً على المحتوى الوارد من الملف [{target_file}] داخل <sources></sources>، "
            "مع ذكر المصطلحات العلمية بالإنجليزية بين قوسين عند ورودها، وبأسلوب واضح وموجز. "
            "إذا كانت الإجابة تتضمن بيانات مجدولة أو سجلات أو مقارنات فاعرضها دائماً في جدول ماركداون منظم بأعمدة واضحة. "
            "إذا لم تجد الإجابة في محتوى هذا الملف تحديداً فقل بصراحة أن المعلومة غير موجودة في هذا الملف."
        )
    else:
        system = (
            "أنت مساعد مذاكرة يجيب بالعربية الفصحى اعتماداً فقط على المحتوى الموجود داخل <sources></sources>، "
            "مع ذكر المصطلحات العلمية بالإنجليزية بين قوسين عند ورودها، وبأسلوب واضح موجز. "
            "إذا كانت الإجابة تتضمن بيانات مجدولة أو سجلات أو مقارنات فاعرضها دائماً في جدول ماركداون منظم بأعمدة واضحة. "
            "إذا لم تجد الإجابة في المحتوى فقل بصراحة أن المعلومة غير موجودة في الملفات."
        )
    prompt = f"<sources>\n{chunks}\n</sources>\n\nالسؤال: {query}"
    payload = {"model": "mada", "messages": [
        {"role": "system", "content": system},
        {"role": "user", "content": prompt},
    ], "temperature": 0.3, "max_tokens": LLM_MAX_TOKENS}
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(
        f"http://{LLM_HOST}:{LLM_PORT}/v1/chat/completions",
        data=data, headers={"Content-Type": "application/json"})
    try:
        with _gen_lock:
            resp = json.loads(urllib.request.urlopen(req, timeout=300).read().decode("utf-8"))
        return resp["choices"][0]["message"]["content"].strip()
    except urllib.error.HTTPError as err:
        try:
            err_body = err.read().decode("utf-8", "replace")
            err_json = json.loads(err_body)
            msg = err_json.get("error", {}).get("message") or err_json.get("message") or str(err)
        except Exception:
            msg = str(err)
        return f"تعذر التوليد ({err.code}): {msg}"
    except Exception as e:
        return f"تعذر التوليد: {e}"


# ── بث الإجابات (وضع الشات، بنمط OpenAI streaming) ──
CHAT_SYSTEM_FREE = (
    "أنت مساعد ذكي محلي يردّ بالعربية الفصحى بأسلوب واضح وسلس، شبيه ChatGPT. "
    "إن سُئلت بالإنجليزية أجب بالإنجليزية، وبالعربية أجب بالعربية. "
    "إذا كانت الإجابة تتضمن بيانات مجدولة فاعرضها في جدول ماركداون منظم. "
    "عند كتابة المعادلات الرياضية أو الرموز العلمية، اكتبها دائماً بصيغة LaTeX القياسية: "
    "للمعادلات المنفصلة استخدم \\[ ... \\] أو $$ ... $$، وللرموز داخل النص استخدم \\( ... \\) أو $ ... $."
)
CHAT_SYSTEM_RAG = (
    "أنت مساعد مراجعة يعتمد فقط على المحتوى داخل <sources></sources> من ملفات المستخدم. "
    "أجب بالعربية الفصحى وأذكر المصطلحات العلمية بالإنجليزية بين قوسين. "
    "إذا كانت الإجابة تتضمن بيانات مجدولة أو سجلات أو مقارنات فاعرضها دائماً في جدول ماركداون منظم بأعمدة واضحة. "
    "عند كتابة المعادلات الرياضية أو الرموز العلمية، اكتبها دائماً بصيغة LaTeX القياسية: "
    "للمعادلات المنفصلة استخدم \\[ ... \\] أو $$ ... $$، وللرموز داخل النص استخدم \\( ... \\) أو $ ... $. "
    "إن لم تكن الإجابة في المحتوى قل بصراحة أن المعلومة غير موجودة في الملفات."
)


def _ensure_llm_ready():
    global _llm_error
    if not ensure_llm():
        return False
    deadline = time.monotonic() + 60
    while time.monotonic() < deadline:
        if _llm_alive():
            return True
        if _llm_proc is not None and _llm_proc.poll() is not None:
            _llm_error = "توقف محرك النموذج (exit=%s)." % _llm_proc.returncode
            return False
        time.sleep(0.5)
    _llm_error = "لم يصبح المحرك جاهزاً خلال 60 ثانية؛ تحقق من الذاكرة والمنفذ 8081."
    return False


def llm_failure_message():
    detail = ""
    try:
        with open(LLM_LOG_PATH, "rb") as log:
            log.seek(0, os.SEEK_END)
            log.seek(max(0, log.tell() - 2000))
            detail = log.read().decode("utf-8", "replace")
        detail = re.sub(r"\x1b\[[0-9;]*m", "", detail)
    except OSError:
        pass
    return (_llm_error or "النموذج المحلي غير متاح") + "\nالسجل: " + LLM_LOG_PATH + ("\n" + detail if detail else "")


# ── البوب-أب السريع (Quick Popup): توليد موجز + محادثة «سريع» ──
QUICK_PROMPTS = {
    "short": "أجب بإجابة صحيحة وموجزة جداً بالعربية في جملة أو جملتين فقط.",
    "medium": "أجب بإجابة صحيحة وواضحة بالعربية في فقرة قصيرة من جملتين إلى ثلاث جمل.",
    "detailed": "أجب بإجابة صحيحة ومفصلة بالعربية مع الشرح الكافي دون إطالة مفرطة.",
}


def quick_system_prompt():
    cfg = load_bg_config()
    cmd = (cfg.get("custom_command") or "").strip()
    if cmd:
        return cmd, 300
    length = cfg.get("length") if cfg.get("length") in QUICK_PROMPTS else "short"
    return QUICK_PROMPTS[length], {"short": 150, "medium": 300, "detailed": 600}.get(length, 150)


def quick_generate(text):
    q = str(text or "").strip()
    if not q:
        return None, "فارغ"
    try:
        system, mt = quick_system_prompt()
        messages = [{"role": "system", "content": system}, {"role": "user", "content": q}]
        parts = list(stream_llm(messages, temperature=0.2, max_tokens=mt))
        ans = "".join(parts).strip()
        if not ans:
            return None, "النموذج لم يُنتج نصاً"
        return ans, None
    except RuntimeError as e:
        return None, str(e)
    except Exception as e:
        return None, "النموذج غير متاح الآن"


def quick_store(q, a):
    now = time.strftime("%Y-%m-%d %H:%M:%S")
    with DB_LOCK:
        conn = get_db()
        r = conn.execute("SELECT * FROM conversations WHERE id = ?", ("quick",)).fetchone()
        if r:
            try:
                msgs = json.loads(r["messages"] or "[]")
            except Exception:
                msgs = []
            msgs.append({"role": "user", "content": q, "ts": now})
            msgs.append({"role": "assistant", "content": a, "ts": now})
            conn.execute(
                "UPDATE conversations SET messages = ?, updated_at = ? WHERE id = ?",
                (json.dumps(msgs, ensure_ascii=False), now, "quick"),
            )
        else:
            msgs = [{"role": "user", "content": q, "ts": now}, {"role": "assistant", "content": a, "ts": now}]
            conn.execute(
                "INSERT INTO conversations (id, mode, title, messages, pinned, archived, created_at, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                ("quick", "free", "سريع", json.dumps(msgs, ensure_ascii=False), 0, 0, now, now),
            )
        conn.commit()
        conn.close()
    return True


# ── إدارة إعدادات النموذج + اختيار المسار عبر نافذة ويندوز ──
def save_model_config(cfg):
    out = load_model_config()
    default_spec = "full" if "cuda" in str(out.get("exe") or LLM_EXE_DEFAULT).lower() else "medium"
    spec = str(cfg.get("spec") or out.get("spec") or default_spec).strip()
    if spec not in SPEC_PARAMS:
        spec = out.get("spec") if out.get("spec") in SPEC_PARAMS else default_spec
    out["spec"] = spec
    for k in ("exe", "model"):
        v = str(cfg.get(k) or "").strip()
        if v:
            out[k] = v
    with open(MODEL_CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
    return out


def pick_file_dialog(title, filter_type="exe", initial_dir=None):
    """يفتح متصفح ملفات ويندوز عبر عملية مساعدة مستقلة دون حظر الخادم، مع مهلة أمان."""
    helper = os.path.join(ROOT, "file_picker_helper.py")
    init_dir = initial_dir if (initial_dir and os.path.isdir(initial_dir)) else ROOT
    if os.path.isfile(helper):
        try:
            cmd = [sys.executable, helper, filter_type, init_dir]
            res = subprocess.run(cmd, capture_output=True, text=True, timeout=50)
            out = (res.stdout or "").strip().splitlines()
            if out:
                candidate = out[-1].strip()
                if candidate and os.path.isfile(candidate):
                    return os.path.normpath(candidate)
        except Exception:
            pass

    # احتياطي مباشر عبر PowerShell STA في عملية منفصلة
    try:
        flt = "Executables (*.exe)|*.exe|All Files (*.*)|*.*" if filter_type == "exe" else "Model Files (*.gguf)|*.gguf|All Files (*.*)|*.*"
        safe_title = (title or "اختر ملف").replace("'", "''")
        safe_dir = init_dir.replace("'", "''")
        ps_code = (
            "$ErrorActionPreference = 'SilentlyContinue'; "
            "[System.Reflection.Assembly]::LoadWithPartialName('System.Windows.Forms') | Out-Null; "
            "$f = New-Object System.Windows.Forms.OpenFileDialog; "
            f"$f.Title = '{safe_title}'; "
            f"$f.Filter = '{flt}'; "
            f"$f.InitialDirectory = '{safe_dir}'; "
            "$f.ShowHelp = $false; "
            "$top = New-Object System.Windows.Forms.Form; "
            "$top.TopMost = $true; "
            "$top.StartPosition = [System.Windows.Forms.FormStartPosition]::CenterScreen; "
            "if ($f.ShowDialog($top) -eq [System.Windows.Forms.DialogResult]::OK) { Write-Output $f.FileName }"
        )
        res = subprocess.run(["powershell", "-NoProfile", "-Sta", "-Command", ps_code],
                             capture_output=True, text=True, timeout=50)
        out = (res.stdout or "").strip()
        if out and os.path.isfile(out):
            return os.path.normpath(out)
    except Exception:
        pass

    return None


def scan_available_models_and_exes():
    """يكتشف تلقائياً ملفات المشغل والنماذج على الجهاز لتسهيل الاختيار بنقرة واحدة."""
    exes = []
    models = []
    c = get_llm_config()
    if c.get("exe") and os.path.isfile(c["exe"]):
        exes.append(os.path.normpath(c["exe"]))
    if c.get("model") and os.path.isfile(c["model"]):
        models.append(os.path.normpath(c["model"]))

    search_dirs = [
        os.path.join(ROOT, "bin"),
        os.path.join(ROOT, "llama"),
        r"D:\jan",
        r"C:\jan",
        os.path.expanduser(r"~\AppData\Local\Programs\Jan"),
        os.path.expanduser(r"~\.cache\lm-studio\models"),
        os.path.expanduser(r"~\.ollama\models"),
    ]

    for base in search_dirs:
        if not os.path.exists(base):
            continue
        try:
            for root_dir, dirs, files in os.walk(base):
                rel = os.path.relpath(root_dir, base)
                if rel.count(os.sep) > 6:
                    dirs.clear()
                    continue
                for f in files:
                    low = f.lower()
                    if low == "llama-server.exe":
                        p = os.path.normpath(os.path.join(root_dir, f))
                        if p not in exes:
                            exes.append(p)
                    elif low.endswith(".gguf"):
                        p = os.path.normpath(os.path.join(root_dir, f))
                        if p not in models:
                            models.append(p)
        except Exception:
            pass

    return {"exes": exes, "models": models}


def handle_fs_ls(query_params):
    """يعيد قائمة المجلدات والملفات لتصفح محلي آمن وسريع داخل المتصفح."""
    path = (query_params.get("path") or [""])[0].strip()
    kind = (query_params.get("kind") or ["all"])[0].strip()

    drives = []
    for letter in "CDEFGHIJKLMNOPQRSTUVWXYZ":
        d = f"{letter}:\\"
        if os.path.exists(d):
            drives.append(d)

    if not path or not os.path.exists(path):
        if kind in ("exe", "model") and os.path.exists("D:\\jan\\llamacpp"):
            path = "D:\\jan\\llamacpp\\models" if kind == "model" else "D:\\jan\\llamacpp\\backends"
        elif os.path.exists("D:\\"):
            path = "D:\\"
        else:
            path = drives[0] if drives else ROOT

    path = os.path.abspath(path)
    entries_dirs = []
    entries_files = []

    try:
        with os.scandir(path) as it:
            for entry in it:
                try:
                    if entry.is_dir(follow_symlinks=False):
                        if not entry.name.startswith((".", "$")):
                            entries_dirs.append(entry.name)
                    elif entry.is_file(follow_symlinks=False):
                        low = entry.name.lower()
                        if kind == "exe" and low.endswith(".exe"):
                            size_mb = round(entry.stat().st_size / (1024 * 1024), 1)
                            entries_files.append({"name": entry.name, "path": os.path.normpath(entry.path), "size": f"{size_mb} MB"})
                        elif kind == "model" and low.endswith(".gguf"):
                            size_gb = round(entry.stat().st_size / (1024 * 1024 * 1024), 2)
                            entries_files.append({"name": entry.name, "path": os.path.normpath(entry.path), "size": f"{size_gb} GB"})
                        elif kind == "all":
                            entries_files.append({"name": entry.name, "path": os.path.normpath(entry.path), "size": ""})
                except Exception:
                    pass
    except Exception as e:
        return {"ok": False, "error": str(e), "current": path, "drives": drives}

    entries_dirs.sort(key=lambda s: s.lower())
    entries_files.sort(key=lambda f: f["name"].lower())

    parent = os.path.dirname(path.rstrip("\\/"))
    if parent and not parent.endswith("\\"):
        parent += "\\"
    if parent == path:
        parent = None

    return {
        "ok": True,
        "current": path,
        "parent": parent,
        "drives": drives,
        "dirs": entries_dirs,
        "files": entries_files
    }


def pick_model_file(kind, current_path=None):
    c = get_llm_config()
    cur = (current_path or "").strip() or c.get("exe" if kind == "exe" else "model", "")
    init_dir = None
    if cur:
        if os.path.isfile(cur):
            init_dir = os.path.dirname(cur)
        elif os.path.isdir(cur):
            init_dir = cur
    if not init_dir:
        try:
            disc = scan_available_models_and_exes()
            target_list = disc.get("exes" if kind == "exe" else "models", [])
            if target_list and os.path.exists(target_list[0]):
                init_dir = os.path.dirname(target_list[0])
        except Exception:
            pass
    title = "اختر برنامج التشغيل llama-server.exe" if kind == "exe" else "اختر ملف النموذج (.gguf)"
    return pick_file_dialog(title, filter_type=kind, initial_dir=init_dir)



# ── المشرف الخلفي (الوكيل الصامت) + خيط الحراسة ──
AGENT_PID = os.path.join(ROOT, ".bgagepid")


def _read_agent_pid():
    try:
        if os.path.isfile(AGENT_PID):
            v = int(open(AGENT_PID, "r").read().strip() or "-1")
            return v if v > 0 else None
    except Exception:
        pass
    return None


def _agent_process_name():
    """اسم العملية الذي يعمل به الوكيل: في النسخة المجمّعة اسم التنفيذي نفسها، وإلا pythonw."""
    if getattr(sys, "frozen", False):
        return os.path.basename(sys.executable).lower()
    return "pythonw"


def _agent_alive():
    """يتحقق من الوكيل عبر كائن العملية وملف PID."""
    global _agent_proc
    if _agent_proc is not None:
        if _agent_proc.poll() is None:
            return True
        _agent_proc = None
    pid = _read_agent_pid()
    if not pid:
        return False
    try:
        import psutil
        if psutil.pid_exists(pid):
            p = psutil.Process(pid)
            if p.is_running() and p.status() != psutil.STATUS_ZOMBIE:
                pname = p.name().lower()
                if "python" in pname or "mada-rag" in pname:
                    return True
    except Exception:
        pass
    try:
        kernel32 = ctypes.windll.kernel32
        h = kernel32.OpenProcess(0x1000, False, pid)
        if h:
            code = ctypes.c_ulong()
            kernel32.GetExitCodeProcess(h, ctypes.byref(code))
            kernel32.CloseHandle(h)
            return code.value == 259
    except Exception:
        pass
    return False


def _kill_stale_agents():
    """يوقف النسخ المتكررة من الوكيل لمنع أي تعارض في الاختصارات."""
    mine = os.getpid()
    try:
        import psutil
        for p in psutil.process_iter(["pid", "name", "cmdline"]):
            try:
                if p.info["pid"] == mine:
                    continue
                cmd = " ".join(p.info["cmdline"] or [])
                pname = (p.info["name"] or "").lower()
                if "--background-agent" in cmd or "background_agent.py" in cmd:
                    p.terminate()
            except Exception:
                pass
    except Exception:
        pass


def start_background_agent():
    global _agent_proc
    cfg = load_bg_config()
    if not cfg.get("enabled", True):
        return
    if _agent_alive():
        return
    _kill_stale_agents()
    script = os.path.join(ROOT, "background_agent.py")
    if not getattr(sys, "frozen", False) and not os.path.isfile(script):
        return
    pyw = os.path.join(os.path.dirname(sys.executable), "pythonw.exe")
    interp = pyw if os.path.isfile(pyw) else sys.executable
    try:
        _agent_proc = subprocess.Popen(
            ([sys.executable, "--background-agent"] if getattr(sys, "frozen", False) else [interp, script]),
            cwd=ROOT, creationflags=subprocess.CREATE_NO_WINDOW
        )
    except Exception:
        _agent_proc = None


def stop_background_agent():
    global _agent_proc
    if _agent_proc is not None:
        try:
            if _agent_proc.poll() is None:
                _agent_proc.terminate()
        except Exception:
            pass
        _agent_proc = None
    _kill_stale_agents()
    try:
        if os.path.isfile(AGENT_PID):
            os.remove(AGENT_PID)
    except Exception:
        pass


def _watchdog_loop():
    while True:
        time.sleep(10)
        try:
            cfg = load_bg_config()
            if not cfg.get("enabled", True):
                continue
            if not _agent_alive():
                start_background_agent()
        except Exception:
            pass


def stream_llm(messages, temperature=0.6, max_tokens=600, endpoint=None, model=None, api_key=None):
    """مولّد: ينشر مقاطع الإجابة من النموذج المحلي أو المخصص يدوياً."""
    parsed_endpoint = urlparse(str(endpoint or ""))
    is_local = (parsed_endpoint.hostname in ("127.0.0.1", "localhost", "::1")
                and parsed_endpoint.port in (8080, LLM_PORT)
                and parsed_endpoint.path.rstrip("/") in ("", "/v1", "/v1/chat/completions"))
    is_custom = bool(endpoint and not is_local)
    if not is_custom:
        if not _ensure_llm_ready():
            raise RuntimeError(llm_failure_message())
        url = f"http://{LLM_HOST}:{LLM_PORT}/v1/chat/completions"
    else:
        url = endpoint.strip().rstrip("/")
        if not url.endswith("/chat/completions"):
            url = f"{url}/chat/completions" if url.endswith("/v1") else f"{url}/v1/chat/completions"

    payload = {
        "model": model or "mada",
        "messages": messages,
        "temperature": float(temperature) if temperature is not None else 0.6,
        "stream": True,
        "max_tokens": int(max_tokens) if max_tokens is not None else 600
    }
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    if api_key and str(api_key).strip():
        k = str(api_key).strip()
        headers["Authorization"] = k if k.startswith("Bearer ") else f"Bearer {k}"

    req = urllib.request.Request(url, data=body, headers=headers)
    try:
        with _gen_lock:
            resp = urllib.request.urlopen(req, timeout=300)
            for raw in resp:
                line = raw.decode("utf-8", "replace").strip()
                if not line.startswith("data:"):
                    continue
                d = line[5:].strip()
                if d == "[DONE]":
                    break
                try:
                    j = json.loads(d)
                except ValueError:
                    continue
                delta = (j.get("choices") or [{}])[0].get("delta", {}).get("content", "")
                if delta:
                    yield delta
    except urllib.error.HTTPError as err:
        try:
            err_body = err.read().decode("utf-8", "replace")
            err_json = json.loads(err_body)
            msg = err_json.get("error", {}).get("message") or err_json.get("message") or str(err)
        except Exception:
            msg = str(err)
        raise RuntimeError(f"خطأ من خادم النموذج ({err.code}): {msg}")


def build_rag_sources(results, max_total_chars=6000):
    lines = []
    total = 0
    for r in results[:8]:
        t = r.get("text", "").strip()[:800]
        a = r.get("ar", "").strip()[:500]
        entry = "[%s]\n%s" % (r.get("file", ""), t)
        if a:
            entry += "\nبالعربية: " + a
        if total + len(entry) > max_total_chars and lines:
            break
        lines.append(entry)
        total += len(entry)
    return "\n\n".join(lines)


DIACRITICS = re.compile(r"[\u064B-\u0652\u0670\u0640\u06D6-\u06ED]")
PUNCT = re.compile(r"[^\w\s\u0600-\u06FF\u0660-\u0669\u066E\u066F\u0640]+", re.UNICODE)


def normalize_ar(s):
    s = DIACRITICS.sub("", s)
    s = s.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    s = s.replace("ة", "ه").replace("ى", "ي").replace("ئ", "ي").replace("ؤ", "و")
    s = s.replace("ـ", "")
    return s


def tokenize(text):
    text = normalize_ar(text).lower()
    text = PUNCT.sub(" ", text)
    words = [w for w in text.split() if w]
    out = list(words)
    for w in words:
        if len(w) >= 4:
            out.append(w[:-2])
            out.append(w[:-1])
    return out


def extract_txt(path):
    txt = open(path, encoding="utf-8", errors="ignore").read()
    return [("main", txt)]


def extract_docx(path):
    d = docx.Document(path)
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    slides = [("وصف", " / ".join(paras))]
    for tbl in d.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                slides.append(("جدول", " | ".join(cells)))
    return slides


def extract_pptx(path):
    p = pptx.Presentation(path)
    slides = []
    for i, sl in enumerate(p.slides, 1):
        parts = []
        for sh in sl.shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    parts.append(t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        if parts:
            slides.append((f"شريحة {i}", "\n".join(parts)))
    return slides


def extract_pdf(path):
    r = PdfReader(path)
    pages = []
    for i, pg in enumerate(r.pages, 1):
        t = pg.extract_text() or ""
        if t.strip():
            pages.append((f"صفحة {i}", t))
    return pages


def extract_xlsx(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sname in wb.sheetnames:
        ws = wb[sname]
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in r if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            sheets.append((f"ورقة {sname}", "\n\n".join(rows)))
    wb.close()
    return sheets


EXTRACTORS = {
    ".txt": extract_txt, ".md": extract_txt, ".sql": extract_txt,
    ".json": extract_txt, ".csv": extract_txt, ".py": extract_txt,
    ".js": extract_txt, ".html": extract_txt, ".css": extract_txt,
    ".log": extract_txt,
    ".docx": extract_docx, ".pptx": extract_pptx, ".pdf": extract_pdf,
    ".xlsx": extract_xlsx,
}


def chunk_text(text, size=700, overlap=90):
    parts = []
    text = re.sub(r"\n{3,}", "\n\n", text.strip())
    para_units = re.split(r"\n\s*\n", text)
    buf = ""
    for pu in para_units:
        pu = pu.strip()
        if not pu:
            continue
        if len(buf) + len(pu) + 2 <= size:
            buf = (buf + "\n\n" + pu).strip()
            continue
        if buf:
            parts.append(buf)
        buf = pu
        while len(buf) > size:
            parts.append(buf[:size])
            buf = buf[size - overlap:]
    if buf:
        parts.append(buf)
    return parts


CHUNKS = []
VECTORIZER = None
MATRIX = None


def rebuild_index():
    global CHUNKS, VECTORIZER, MATRIX
    CHUNKS = []
    for fname in sorted(os.listdir(UPLOADS)):
        fpath = os.path.join(UPLOADS, fname)
        ext = os.path.splitext(fname)[1].lower()
        if ext not in EXTRACTORS or not os.path.isfile(fpath):
            continue
        try:
            sections = EXTRACTORS[ext](fpath)
        except Exception as e:
            print("extract fail", fname, e)
            continue
        for ref, text in sections:
            for i, c in enumerate(chunk_text(text), 1):
                CHUNKS.append({"f": fname, "ref": ref, "i": i, "t": c})
    texts = [c["t"] for c in CHUNKS]
    if texts:
        VECTORIZER = TfidfVectorizer(tokenizer=tokenize, lowercase=False, sublinear_tf=True, min_df=1)
        MATRIX = VECTORIZER.fit_transform(texts)
    else:
        VECTORIZER = None
        MATRIX = None
    with open(CHUNKS_PATH, "w", encoding="utf-8") as f:
        json.dump(CHUNKS, f, ensure_ascii=False)
    print("index rebuilt:", len(CHUNKS), "chunks from", len(os.listdir(UPLOADS)))


def save_chunks(chunks):
    with open(CHUNKS_PATH, "w", encoding="utf-8") as f:
        json.dump(chunks, f, ensure_ascii=False)


def _normalize_target_files(target_file=None, target_files=None):
    if target_files:
        if isinstance(target_files, (list, tuple, set)):
            s = {os.path.basename(str(f)).strip() for f in target_files if f and str(f).strip()}
            return s if s else None
        return {os.path.basename(str(target_files)).strip()}
    if target_file and str(target_file).strip():
        return {os.path.basename(str(target_file)).strip()}
    return None


def search(q, top_n=6, target_file=None, target_files=None):
    t_files = _normalize_target_files(target_file, target_files)
    if VECTORIZER is None or MATRIX is None or not CHUNKS:
        return []
    qv = VECTORIZER.transform([q])
    scores = (MATRIX @ qv.T).toarray().ravel()
    qtoks = set(tokenize(q))
    for i, sc in enumerate(scores):
        if t_files and CHUNKS[i]["f"] not in t_files:
            scores[i] = 0
            continue
        wtoks = set(tokenize(CHUNKS[i]["t"]))
        inter = len(qtoks & wtoks)
        if inter:
            scores[i] += 0.25 * (inter / max(1, len(qtoks)))
    order = scores.argsort()[::-1]
    out = []
    seen = set()
    for i in order:
        if not scores[i] or scores[i] <= 0:
            continue
        if t_files and CHUNKS[i]["f"] not in t_files:
            continue
        key = (CHUNKS[i]["f"], CHUNKS[i]["ref"], CHUNKS[i]["i"] // 2)
        if key in seen:
            continue
        seen.add(key)
        c = CHUNKS[i]
        out.append({
            "text": c["t"], "file": c["f"], "ref": c["ref"],
            "score": round(float(scores[i]) * 100, 1),
        })
        if len(out) >= top_n:
            break
    return out


def _scores(q, target_file=None, target_files=None):
    t_files = _normalize_target_files(target_file, target_files)
    if VECTORIZER is None or MATRIX is None or not CHUNKS:
        return []
    qv = VECTORIZER.transform([q])
    scores = (MATRIX @ qv.T).toarray().ravel()
    qtoks = set(tokenize(q))
    for i, sc in enumerate(scores):
        if t_files and CHUNKS[i]["f"] not in t_files:
            continue
        wtoks = set(tokenize(CHUNKS[i]["t"]))
        inter = len(qtoks & wtoks)
        if inter:
            scores[i] += 0.25 * (inter / max(1, len(qtoks)))
    return [(float(v), i) for i, v in enumerate(scores) if v > 0 and (not t_files or CHUNKS[i]["f"] in t_files)]


def has_arabic(s):
    return any("\u0600" <= c <= "\u06FF" for c in s)


def translate_query(q):
    if not has_arabic(q) or not translation_available():
        return ""
    return translate_text(q, "ar2en") or ""


def search_dual(q, top_n=6, target_file=None, target_files=None):
    """إن كان السؤال عربياً والملفات إنجليزية: نترجم السؤال ونبحث به أيضاً."""
    t_files = _normalize_target_files(target_file, target_files)
    if not has_arabic(q):
        return search(q, top_n, target_files=t_files)
    merged = {}
    for sc, i in _scores(q, target_files=t_files):
        if sc > (merged.get(i, (0, 0))[0]):
            merged[i] = (sc, i)
    qen = translate_query(q).strip()
    if qen:
        for sc, i in _scores(qen, target_files=t_files):
            if sc > (merged.get(i, (0, 0))[0]):
                merged[i] = (sc, i)
    order = sorted(merged.values(), key=lambda x: -x[0])
    out = []
    seen = set()
    for sc, i in order:
        if t_files and CHUNKS[i]["f"] not in t_files:
            continue
        key = (CHUNKS[i]["f"], CHUNKS[i]["ref"], CHUNKS[i]["i"] // 2)
        if key in seen:
            continue
        seen.add(key)
        c = CHUNKS[i]
        out.append({
            "text": c["t"], "file": c["f"], "ref": c["ref"],
            "score": round(sc * 100, 1),
        })
        if len(out) >= top_n:
            break
    return out


class ActiveGeneration:
    """إدارة توليد خلفي مستقل للمحادثة لضمان استمرار التوليد حتى لو انقطع الاتصال أو أُعيد تحميل الصفحة."""
    def __init__(self, chat_id, mode, is_temp=False):
        self.chat_id = chat_id
        self.mode = mode
        self.is_temp = is_temp
        self.full_text = ""
        self.meta = None
        self.events = []
        self.subscribers = []
        self.lock = threading.Lock()
        self.done = False
        self.error = None
        self.stop_requested = False
        self.created_at = time.time()
        self.completed_at = None

    def add_subscriber(self):
        with self.lock:
            q = queue.Queue()
            for ev in self.events:
                q.put(ev)
            if self.done:
                q.put(None)
            else:
                self.subscribers.append(q)
            return q

    def remove_subscriber(self, q):
        with self.lock:
            if q in self.subscribers:
                self.subscribers.remove(q)

    def emit(self, ev):
        with self.lock:
            self.events.append(ev)
            if ev.get("type") == "txt":
                self.full_text += ev.get("t", "")
            elif ev.get("type") == "meta":
                self.meta = ev
            elif ev.get("type") == "err":
                self.error = ev.get("msg")
            for q in list(self.subscribers):
                try:
                    q.put(ev)
                except Exception:
                    pass

    def finish(self, err_msg=None):
        with self.lock:
            if self.done:
                return
            self.done = True
            self.completed_at = time.time()
            if err_msg:
                self.error = err_msg
                ev = {"type": "err", "msg": err_msg}
                self.events.append(ev)
                for q in list(self.subscribers):
                    try:
                        q.put(ev)
                    except Exception:
                        pass
            ev_done = {"type": "done"}
            self.events.append(ev_done)
            for q in list(self.subscribers):
                try:
                    q.put(ev_done)
                    q.put(None)
                except Exception:
                    pass
            self.subscribers.clear()

        # حفظ تلقائي للرد المكتمل في SQLite إذا لم تكن محادثة مؤقتة
        if not self.is_temp and self.chat_id and self.full_text.strip():
            try:
                conv = get_conv(self.chat_id)
                if conv is None:
                    conv = {
                        "id": self.chat_id,
                        "mode": self.mode,
                        "title": self.full_text.replace("\n", " ").strip()[:42] or "محادثة",
                        "messages": [],
                    }
                if conv:
                    msgs = conv.get("messages") or []
                    if msgs and msgs[-1].get("role") == "assistant":
                        msgs[-1]["content"] = self.full_text
                    else:
                        msgs.append({"role": "assistant", "content": self.full_text})
                    conv["messages"] = msgs
                    upsert_conv(conv)
            except Exception as e:
                print(f"[ActiveGen] Error auto-saving to SQLite: {e}")


ACTIVE_GENERATIONS = {}
GEN_MGR_LOCK = threading.Lock()


def get_or_create_generation(chat_id, mode, is_temp=False):
    with GEN_MGR_LOCK:
        now = time.time()
        for cid in list(ACTIVE_GENERATIONS.keys()):
            g = ACTIVE_GENERATIONS[cid]
            if g.done and (now - (g.completed_at or g.created_at) > 300):
                ACTIVE_GENERATIONS.pop(cid, None)

        gen = ACTIVE_GENERATIONS.get(chat_id)
        if gen and not gen.done:
            return gen, False
        gen = ActiveGeneration(chat_id, mode, is_temp)
        ACTIVE_GENERATIONS[chat_id] = gen
        return gen, True


def get_active_generation(chat_id):
    with GEN_MGR_LOCK:
        return ACTIVE_GENERATIONS.get(chat_id)


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def _send(self, code, ctype, body):
        try:
            self.send_response(code)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            self.wfile.write(body)
        except (ConnectionAbortedError, ConnectionResetError, BrokenPipeError):
            self.close_connection = True

    def _json(self, obj, code=200):
        self._send(code, "application/json; charset=utf-8",
                   json.dumps(obj, ensure_ascii=False).encode("utf-8"))

    def _sse(self, events):
        self.send_response(200)
        self.send_header("Content-Type", "text/event-stream; charset=utf-8")
        self.send_header("Cache-Control", "no-cache")
        self.send_header("Connection", "keep-alive")
        self.end_headers()
        for e in events:
            try:
                self.wfile.write(("data: " + json.dumps(e, ensure_ascii=False) + "\n\n").encode("utf-8"))
                self.wfile.flush()
            except Exception:
                break

    def _sse_stream_gen(self, gen):
        q = gen.add_subscriber()
        self.send_response(200)
        self.send_header("Content-Type", "text/event-stream; charset=utf-8")
        self.send_header("Cache-Control", "no-cache")
        self.send_header("Connection", "keep-alive")
        self.end_headers()
        try:
            while True:
                try:
                    ev = q.get(timeout=20)
                    if ev is None:
                        break
                    self.wfile.write(("data: " + json.dumps(ev, ensure_ascii=False) + "\n\n").encode("utf-8"))
                    self.wfile.flush()
                except queue.Empty:
                    if gen.done:
                        break
                    try:
                        self.wfile.write(b": keep-alive\n\n")
                        self.wfile.flush()
                    except Exception:
                        break
        except Exception:
            pass
        finally:
            gen.remove_subscriber(q)

    def _events_free(self, messages, temperature=0.6, max_tokens=600, endpoint=None, model=None, api_key=None, memory_enabled=True):
        try:
            block = memory_rules_block() if memory_enabled else ""
            if block:
                if messages and messages[0].get("role") == "system":
                    messages = [dict(messages[0], content=messages[0]["content"] + "\n\n" + block)] + messages[1:]
                else:
                    messages = [{"role": "system", "content": block}] + messages
            for chunk in stream_llm(messages, temperature=temperature, max_tokens=max_tokens, endpoint=endpoint, model=model, api_key=api_key):
                yield {"type": "txt", "t": chunk}
        except RuntimeError as ex:
            yield {"type": "err", "msg": str(ex)}
            return
        except Exception as ex:
            yield {"type": "err", "msg": f"فشل أثناء التوليد: {ex}"}
            return
        yield {"type": "done"}

    def _events_rag(self, q, translate, target_file=None, target_files=None, temperature=0.3, max_tokens=600, endpoint=None, model=None, api_key=None, custom_instructions=None, memory_enabled=True):
        t_files = _normalize_target_files(target_file, target_files)
        try:
            with LOCK:
                res = search_dual(q, target_files=t_files)
            if not res:
                yield {"type": "done"}
                return
            if translate:
                with LOCK:
                    res = translate_results(res)
            files = sorted({r["file"] for r in res[:8]})
            yield {"type": "meta", "sources": len(res), "files": files}
            if t_files:
                flist_str = "، ".join(sorted(t_files))
                system_prompt = (
                    f"أنت مساعد مراجعة يعتمد فقط وحصرياً على المحتوى من ملفات [{flist_str}] داخل <sources></sources>. "
                    "أجب بالعربية الفصحى واذكر المصطلحات العلمية بالإنجليزية بين قوسين. "
                    "إذا لم تجد الإجابة في محتوى هذه الملفات المحددة تحديداً فقل بصراحة أن المعلومة غير موجودة في هذه الملفات."
                )
            else:
                system_prompt = CHAT_SYSTEM_RAG
            if custom_instructions and str(custom_instructions).strip():
                system_prompt += f"\n\n[تعليمات وتفضيلات مخصصة من المستخدم]:\n{custom_instructions}"
            block = memory_rules_block() if memory_enabled else ""
            if block:
                system_prompt += "\n\n" + block
            msgs = [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": "<sources>\n%s\n</sources>\n\nالسؤال: %s" % (build_rag_sources(res), q)},
            ]
            for chunk in stream_llm(msgs, temperature=temperature, max_tokens=max_tokens, endpoint=endpoint, model=model, api_key=api_key):
                yield {"type": "txt", "t": chunk}
        except RuntimeError as ex:
            yield {"type": "err", "msg": str(ex)}
            return
        except Exception as ex:
            yield {"type": "err", "msg": f"فشل أثناء التوليد: {ex}"}
            return
        yield {"type": "done"}

    def do_GET(self):
        u = urlparse(self.path)
        if u.path in ("/", "/index.html"):
            index_path = get_asset_path("index.html")
            if os.path.isfile(index_path):
                with open(index_path, "rb") as f:
                    self._send(200, "text/html; charset=utf-8", f.read())
            else:
                self._send(404, "text/plain; charset=utf-8", "واجهة البحث غير متوفرة".encode("utf-8"))
        elif u.path == "/chat":
            chat_path = get_asset_path("chat.html")
            if os.path.isfile(chat_path):
                with open(chat_path, "rb") as f:
                    self._send(200, "text/html; charset=utf-8", f.read())
            else:
                self._send(404, "text/plain; charset=utf-8", "صفحة المحادثة غير متوفرة".encode("utf-8"))
        elif u.path.startswith("/static/"):
            rel_path = u.path[1:].replace("/", os.sep)
            local_path = get_asset_path(rel_path)
            if os.path.isfile(local_path):
                ctype = "application/octet-stream"
                if local_path.endswith(".js"): ctype = "application/javascript; charset=utf-8"
                elif local_path.endswith(".css"): ctype = "text/css; charset=utf-8"
                elif local_path.endswith(".woff2"): ctype = "font/woff2"
                elif local_path.endswith(".woff"): ctype = "font/woff"
                elif local_path.endswith(".ttf"): ctype = "font/ttf"
                with open(local_path, "rb") as f:
                    self._send(200, ctype, f.read())
            else:
                self._send(404, "text/plain", b"File not found")
        elif u.path == "/docs":
            raw_files = sorted(
                f for f in os.listdir(UPLOADS)
                if os.path.splitext(f)[1].lower() in ALLOWED
            )
            files_meta = []
            for fname in raw_files:
                fpath = os.path.join(UPLOADS, fname)
                sz = os.path.getsize(fpath) if os.path.isfile(fpath) else 0
                chunks_cnt = sum(1 for c in CHUNKS if c.get("f") == fname)
                files_meta.append({
                    "name": fname,
                    "size": sz,
                    "chunks": chunks_cnt,
                    "ext": os.path.splitext(fname)[1].lower()
                })
            n = len(CHUNKS)
            words = sum(len(c["t"].split()) for c in CHUNKS)
            self._json({
                "files": raw_files,
                "files_meta": files_meta,
                "chunks": n,
                "words": words,
                "translate": translation_available(),
                "gen": generation_available(),
                "v": 5
            })
        elif u.path == "/ask":
            qp = parse_qs(u.query)
            q = unquote(qp.get("q", [""])[0])
            target_file = unquote(qp.get("file", [""])[0]).strip() or None
            tr = qp.get("tr", ["0"])[0] == "1"
            gen = qp.get("gen", ["0"])[0] == "1"
            with LOCK:
                res = search_dual(q, target_file=target_file)
            if tr and res:
                res = translate_results(res)
            gen_answer = None
            if gen and res:
                gen_answer = generate_answer(q, res, target_file=target_file)
            self._json({"query": q, "file": target_file, "results": res, "generated": gen_answer})
        elif u.path == "/convs":
            qp = parse_qs(u.query)
            archived = qp.get("archived", ["0"])[0] == "1"
            c_list = list_convs(include_archived=archived)
            self._json({"ok": True, "conversations": c_list, "chats": c_list})
        elif u.path in ("/projects", "/projects/"):
            accept = self.headers.get("Accept", "")
            if "text/html" in accept:
                chat_path = get_asset_path("chat.html")
                if os.path.isfile(chat_path):
                    with open(chat_path, "rb") as f:
                        self._send(200, "text/html; charset=utf-8", f.read())
                else:
                    self._send(404, "text/plain; charset=utf-8", "صفحة المحادثة غير متوفرة".encode("utf-8"))
            else:
                self._json({"ok": True, "projects": list_projects()})
        elif u.path in ("/project", "/project/") or u.path.startswith("/project/"):
            chat_path = get_asset_path("chat.html")
            if os.path.isfile(chat_path):
                with open(chat_path, "rb") as f:
                    self._send(200, "text/html; charset=utf-8", f.read())
            else:
                self._send(404, "text/plain; charset=utf-8", "صفحة المحادثة غير متوفرة".encode("utf-8"))
        elif u.path == "/api/projects":
            self._json({"ok": True, "projects": list_projects()})
        elif u.path == "/api/generation/status":
            qs = parse_qs(u.query)
            cid = (qs.get("chat_id") or [""])[0].strip()
            gen = get_active_generation(cid)
            if gen and not gen.done:
                self._json({
                    "ok": True,
                    "active": True,
                    "status": "running",
                    "chat_id": cid,
                    "text": gen.full_text,
                    "meta": gen.meta,
                    "done": False
                })
            elif gen and gen.done:
                self._json({
                    "ok": True,
                    "active": False,
                    "status": "completed",
                    "chat_id": cid,
                    "text": gen.full_text,
                    "meta": gen.meta,
                    "done": True
                })
            else:
                self._json({"ok": True, "active": False, "status": "none", "chat_id": cid})
        elif u.path == "/api/generation/stream":
            qs = parse_qs(u.query)
            cid = (qs.get("chat_id") or [""])[0].strip()
            gen = get_active_generation(cid)
            if gen:
                self._sse_stream_gen(gen)
            else:
                self._json({"ok": False, "error": "no active generation"}, 404)
        elif u.path == "/api/generation/active":
            with GEN_MGR_LOCK:
                active = [{
                    "chat_id": g.chat_id,
                    "is_temp": g.is_temp,
                    "mode": g.mode,
                    "text": g.full_text,
                    "meta": g.meta
                } for g in ACTIVE_GENERATIONS.values() if not g.done]
            self._json({"ok": True, "generations": active})
        elif u.path == "/api/memory":
            self._json({"ok": True, "rules": list_memory()})
        elif u.path == "/api/bg":
            self._json({"ok": True, "settings": load_bg_config()})
        elif u.path == "/api/model":
            self._json({"ok": True, "model": get_llm_config(), "discovered": scan_available_models_and_exes()})
        elif u.path == "/api/model/scan":
            self._json({"ok": True, "discovered": scan_available_models_and_exes()})
        elif u.path == "/api/fs/ls":
            qs = parse_qs(u.query)
            self._json(handle_fs_ls(qs))
        else:
            self._send(404, "text/plain", b"not found")

    def do_POST(self):
        u = urlparse(self.path)
        length = int(self.headers.get("Content-Length", 0))
        body = self.rfile.read(length)
        if u.path == "/api/memory":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            action = str(data.get("action") or "add")
            kind = str(data.get("kind") or "term")
            if kind not in ("term", "rule"):
                kind = "term"
            subject = str(data.get("subject") or "").strip()
            wrong = str(data.get("wrong") or "").strip()
            note = str(data.get("note") or "").strip()
            if action == "delete":
                mid = data.get("id")
                if mid is None:
                    self._json({"ok": False, "error": "id required"}, 400)
                    return
                delete_memory(int(mid))
                self._json({"ok": True, "deleted": int(mid)})
            elif action == "update":
                mid = data.get("id")
                if mid is None:
                    self._json({"ok": False, "error": "id required"}, 400)
                    return
                if not subject:
                    self._json({"ok": False, "error": "subject required"}, 400)
                    return
                update_memory(int(mid), kind, subject, wrong, note)
                self._json({"ok": True, "id": int(mid)})
            else:
                if not subject:
                    self._json({"ok": False, "error": "subject required"}, 400)
                    return
                new_id = add_memory(kind, subject, wrong, note)
                self._json({"ok": True, "id": new_id})
        elif u.path == "/api/bg":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            saved = save_bg_config(sanitize_bg(data))
            self._json({"ok": True, "settings": saved})
        elif u.path == "/api/model":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            before = get_llm_config()
            saved = save_model_config(data)
            after = get_llm_config()
            changed = (before["exe"], before["model"], before["spec"]) != (
                after["exe"], after["model"], after["spec"])
            if changed:
                stop_llm()
                threading.Thread(target=ensure_llm, daemon=True).start()
            self._json({"ok": True, "model": after, "restarted": changed})
        elif u.path == "/api/model/pick":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            kind = str(data.get("kind") or "model")
            current_path = str(data.get("current_path") or "").strip()
            path = pick_model_file(kind if kind in ("exe", "model") else "model", current_path=current_path)
            self._json({"ok": True, "path": path})
        elif u.path == "/api/quick":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            text = str(data.get("text") or "").strip()
            if not text:
                self._json({"ok": False, "error": "نص فارغ"}, 400)
                return
            answer, err = quick_generate(text)
            if answer is None:
                self._json({"ok": False, "error": err or "تعذر توليد الرد"}, 500)
                return
            cfg = load_bg_config()
            if cfg.get("saveReplies", True):
                quick_store(text, answer)
            self._json({"ok": True, "answer": answer})
        elif u.path == "/api/generation/stop":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            cid = str(data.get("chat_id") or "").strip()
            gen = get_active_generation(cid)
            if gen and not gen.done:
                gen.stop_requested = True
                gen.finish()
                self._json({"ok": True, "stopped": True, "text": gen.full_text})
            else:
                self._json({"ok": True, "stopped": False})
        elif u.path == "/cgen":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            chat_id = str(data.get("chat_id") or "").strip()
            is_temp = bool(data.get("is_temp", False))
            reconnect = bool(data.get("reconnect", False))

            if chat_id and reconnect:
                gen = get_active_generation(chat_id)
                if gen:
                    self._sse_stream_gen(gen)
                    return

            messages = [{"role": "user", "content": data["q"]}] if data.get("q") else []
            if data.get("messages"):
                messages = data["messages"]
            if not messages or not messages[-1].get("content"):
                self._json({"ok": False, "error": "empty"})
                return
            msgs = [{"role": m["role"], "content": str(m["content"])[:6000]} for m in messages]
            if msgs[0].get("role") != "system":
                msgs.insert(0, {"role": "system", "content": CHAT_SYSTEM_FREE})
            
            temp = float(data.get("temperature", 0.6)) if data.get("temperature") is not None else 0.6
            max_t = int(data.get("max_tokens", 600)) if data.get("max_tokens") is not None else 600
            endpoint = data.get("endpoint")
            model_name = data.get("model")
            api_key = data.get("api_key")
            mem_en = data.get("memory_enabled", True)
            if not isinstance(mem_en, bool):
                mem_en = str(mem_en).lower() != "false"

            if chat_id:
                gen, is_new = get_or_create_generation(chat_id, "free", is_temp)
                if is_new:
                    def _worker_free():
                        try:
                            for ev in self._events_free(msgs, temperature=temp, max_tokens=max_t,
                                                       endpoint=endpoint, model=model_name, api_key=api_key,
                                                       memory_enabled=mem_en):
                                if gen.stop_requested:
                                    break
                                gen.emit(ev)
                        except Exception as ex:
                            gen.finish(str(ex))
                        finally:
                            gen.finish()
                    threading.Thread(target=_worker_free, daemon=True).start()
                self._sse_stream_gen(gen)
            else:
                self._sse(self._events_free(msgs, temperature=temp, max_tokens=max_t, endpoint=endpoint, model=model_name, api_key=api_key, memory_enabled=mem_en))
        elif u.path == "/ragen":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            chat_id = str(data.get("chat_id") or "").strip()
            is_temp = bool(data.get("is_temp", False))
            reconnect = bool(data.get("reconnect", False))

            if chat_id and reconnect:
                gen = get_active_generation(chat_id)
                if gen:
                    self._sse_stream_gen(gen)
                    return

            q = (data.get("q") or "").strip()
            if not q:
                self._json({"ok": False, "error": "empty"})
                return
            tr = bool(data.get("tr", True))
            t_file = (data.get("file") or "").strip() or None
            t_files = data.get("files") or None
            temp = float(data.get("temperature", 0.3)) if data.get("temperature") is not None else 0.3
            max_t = int(data.get("max_tokens", 800)) if data.get("max_tokens") is not None else 800
            endpoint = data.get("endpoint")
            model_name = data.get("model")
            api_key = data.get("api_key")
            custom_inst = data.get("custom_instructions")
            mem_en = data.get("memory_enabled", True)
            if not isinstance(mem_en, bool):
                mem_en = str(mem_en).lower() != "false"

            if chat_id:
                gen, is_new = get_or_create_generation(chat_id, "review", is_temp)
                if is_new:
                    def _worker_rag():
                        try:
                            for ev in self._events_rag(q, tr, target_file=t_file, target_files=t_files,
                                                       temperature=temp, max_tokens=max_t,
                                                       endpoint=endpoint, model=model_name, api_key=api_key,
                                                       custom_instructions=custom_inst,
                                                       memory_enabled=mem_en):
                                if gen.stop_requested:
                                    break
                                gen.emit(ev)
                        except Exception as ex:
                            gen.finish(str(ex))
                        finally:
                            gen.finish()
                    threading.Thread(target=_worker_rag, daemon=True).start()
                self._sse_stream_gen(gen)
            else:
                self._sse(self._events_rag(q, tr, target_file=t_file, target_files=t_files,
                                           temperature=temp, max_tokens=max_t,
                                           endpoint=endpoint, model=model_name, api_key=api_key,
                                           custom_instructions=custom_inst,
                                           memory_enabled=mem_en))
        elif u.path == "/delete":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            fname = (data.get("file") or "").strip()
            if not fname:
                self._json({"ok": False, "error": "اسم الملف مطلوب"}, 400)
                return
            safe_fname = os.path.basename(fname)
            fpath = os.path.join(UPLOADS, safe_fname)
            if not os.path.abspath(fpath).startswith(os.path.abspath(UPLOADS)):
                self._json({"ok": False, "error": "مسار غير مسموح به"}, 403)
                return
            if not os.path.isfile(fpath):
                self._json({"ok": False, "error": "الملف غير موجود"}, 404)
                return
            try:
                os.remove(fpath)
            except Exception as e:
                self._json({"ok": False, "error": f"تعذر حذف الملف: {e}"}, 500)
                return
            with LOCK:
                rebuild_index()
            self._json({"ok": True, "file": safe_fname, "chunks": len(CHUNKS)})
            return
        elif u.path == "/convs":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            chat = data.get("chat") or data
            ok = upsert_conv(chat)
            self._json({"ok": ok})
            return
        elif u.path == "/convs/delete":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            cid = data.get("id")
            self._json({"ok": delete_conv(cid) if cid else False})
            return
        elif u.path == "/convs/move":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            cid = data.get("id")
            pid = data.get("project_id")
            self._json({"ok": move_conv(cid, pid) if cid else False})
            return
        elif u.path == "/convs/pin":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            cid = data.get("id")
            pinned = data.get("pinned", True)
            self._json({"ok": pin_conv(cid, pinned) if cid else False})
            return
        elif u.path == "/convs/archive":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            cid = data.get("id")
            archived = data.get("archived", True)
            self._json({"ok": archive_conv(cid, archived) if cid else False})
            return
        elif u.path == "/convs/clear":
            self._json({"ok": clear_convs()})
            return
        elif u.path == "/projects":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            res = upsert_project(data)
            self._json({"ok": True, "project": res})
            return
        elif u.path == "/projects/update":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            res = upsert_project(data)
            self._json({"ok": True, "project": res})
            return
        elif u.path == "/projects/delete":
            try:
                data = json.loads(body.decode("utf-8") or "{}")
            except ValueError:
                data = {}
            pid = data.get("id")
            self._json({"ok": delete_project(pid) if pid else False})
            return
        elif u.path != "/upload":
            self._json({"ok": False, "error": "bad path"}, 404)
            return
        else:
            ctype = self.headers.get("Content-Type", "")
            m = re.search(r"boundary=(.+)", ctype)
            if not m:
                self._json({"ok": False, "error": "no boundary"})
                return
            boundary = m.group(1).strip().strip('"').encode()
            saved, skipped = [], []
            for part in body.split(b"--" + boundary):
                part = part.strip(b"\r\n")
                if not part or part == b"--":
                    continue
                head, _, content = part.partition(b"\r\n\r\n")
                disp = ""
                for line in head.decode("utf-8", "ignore").split("\r\n"):
                    if line.lower().startswith("content-disposition"):
                        disp = line
                fm = re.search(r'filename="([^"]*)"', disp)
                if not fm:
                    continue
                fname = os.path.basename(fm.group(1))
                ext = os.path.splitext(fname)[1].lower()
                if ext not in ALLOWED:
                    skipped.append(fname + " (صيغة غير مدعومة)")
                    continue
                if not content:
                    skipped.append(fname + " (فارغ)")
                    continue
                path = os.path.join(UPLOADS, fname)
                n = 1
                base, e = os.path.splitext(fname)
                while os.path.exists(path):
                    path = os.path.join(UPLOADS, f"{base}_{n}{e}")
                    n += 1
                with open(path, "wb") as f:
                    f.write(content)
                saved.append(os.path.basename(path))
            with LOCK:
                rebuild_index()
            self._json({"ok": True, "saved": saved, "skipped": skipped,
                        "chunks": len(CHUNKS)})


def main():
    init_db()
    os.makedirs(UPLOADS, exist_ok=True)
    os.makedirs(os.path.dirname(CHUNKS_PATH), exist_ok=True)
    with LOCK:
        rebuild_index()

    def _warm():
        try:
            if ensure_bridge():
                translate_text("Warm-up sentence: mitochondria release energy as ATP.")
        except Exception:
            pass
    if translation_available():
        threading.Thread(target=_warm, daemon=True).start()
    port = 8787
    srv = ThreadingHTTPServer(("127.0.0.1", port), Handler)
    trs = "متاحة (نموذجك المحلي إنجليزي-عربي)" if translation_available() else "غير متاحة (عدّل مسار VENV في السيرفر)"
    gens = "متاح (نموذج محلي)" if generation_available() else "غير مثبّت"
    print(f"مذكّرتي تعمل: http://127.0.0.1:{port}")
    print(f"الترجمة المحلية: {trs}")
    print(f"التوليد المحلي: {gens}")
    threading.Thread(target=_watchdog_loop, daemon=True).start()
    start_background_agent()
    try:
        import webbrowser
        threading.Timer(0.7, lambda: webbrowser.open(f"http://127.0.0.1:{port}")).start()
    except Exception:
        pass
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        stop_background_agent()
        print("bye")


if __name__ == "__main__":
    main()
